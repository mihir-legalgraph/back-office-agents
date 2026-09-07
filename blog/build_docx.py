"""blog.md -> the publication .docx, in the house blog format.

Run: python build_docx.py
Deliberately a small markdown subset - headings, bullets, numbered steps, tables,
one pull-quote, inline bold/italic/code. pandoc mangles the tables, so this exists.
"""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "blog.md"
OUT = "Solving Contract Lookup Bottlenecks Using a Claude Skill.docx"
HERO = "../legal/agiloft-skill/assets/architecture.png"
HERO_CAPTION = ("Solid lines run every time. Dashed lines are conditional - the refetch loop fires "
                "only when a page comes back full, and the Agiloft widget figure appears only when "
                "it differs from ours.")

# --- inline runs -------------------------------------------------------------
TOKEN = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`[^`]+`)", re.S)

def runs(par, text, bold=False, italic=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        b, i, code = bold, italic, False
        if tok.startswith("***") and tok.endswith("***"):
            tok, b, i = tok[3:-3], True, True
        elif tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, i = tok[1:-1], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, code = tok[1:-1], True
        r = par.add_run(tok)
        r.bold, r.italic = b, i
        if code:
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
    return par

def para(doc, text, style=None, **kw):
    p = doc.add_paragraph(style=style)
    runs(p, text, **kw)
    return p

# --- table -------------------------------------------------------------------
def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], cells[2:]          # cells[1] is the |---| separator
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, txt in zip(t.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        runs(cell.paragraphs[0], txt, bold=True)
    for row in body:
        for cell, txt in zip(t.add_row().cells, row):
            cell.paragraphs[0].text = ""
            runs(cell.paragraphs[0], txt)
    doc.add_paragraph()

# --- document ----------------------------------------------------------------
def build():
    text = open(SRC, encoding="utf-8").read().replace(" — ", " - ").replace("—", "-")
    lines = text.split("\n")
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i, first_para, hero_done = 0, True, False
    while i < len(lines):
        ln = lines[i].rstrip()
        i += 1
        if not ln.strip():
            continue

        if ln.startswith("|"):                                   # table
            block = [ln]
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            add_table(doc, block)

        elif ln.startswith("#"):                                 # heading
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("# ").strip(), level=min(level, 3))
            if ln.lstrip("# ").strip() == "Architecture" and not hero_done:
                para(doc, "Here is a snapshot of how the solution works:")
                doc.add_picture(HERO, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = para(doc, HERO_CAPTION, style="Caption", italic=True)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                hero_done = True

        elif ln.startswith("> "):                                # pull quote
            p = para(doc, ln[2:].strip(), style="Intense Quote")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        elif ln.startswith("- "):                                # bullet
            para(doc, ln[2:].strip(), style="List Bullet")

        elif re.match(r"^\d+\.\s", ln):                          # numbered step
            para(doc, re.sub(r"^\d+\.\s", "", ln), style="List Number")

        elif first_para:                                         # kicker + byline
            p = para(doc, ln.strip("*"), italic=True)
            p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            if "MIN READ" in ln:
                continue
            first_para = False
        else:
            para(doc, ln.strip())

    doc.save(OUT)
    print("wrote", OUT)

if __name__ == "__main__":
    try:
        build()
    except PermissionError:
        sys.exit(f"{OUT} is open in Word - close it and re-run.")
